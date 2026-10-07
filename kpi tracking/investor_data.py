"""
Investor Data — fetches previous-month figures for the "Investor Data" Google Sheet
(one tab per market: SG, VN, KH, TH, HK, NYC, CO), writes a clean Excel and uploads
it to Slack. Same style as regional_od.py.

Rows per market (same order as the Google Sheet):
    Completed Trips, GMV, Platform Revenue,
    Rider  - Registered (Cumulative Sign Up), Active (MAU), Completed,
    Driver - Registered (Cumulative Sign Up), Active (MAU), Completed

Sources (validated against the sheet for Aug + Sep 2026 — see investor_data_validation.md):
    Completed / GMV / Platform Revenue reuse the regional_od.py queries where the sheet
    uses the same figure (2183 / 3771 / 3106 / 7579 / 6561 / 6189 / 7644 / 7655).
    VN and KH GMV use 6008 (car-type-restricted "Regional - GMV Trend"), KH completed
    trips use the KH perf-sheet query 1625 — those are the figures the sheet carries.
    Cumulative sign-ups: 8095 (riders) and 7913 (drivers), one call for every market.
    MAU / completed users come from the monthly perf-sheet queries each market already
    runs in monthly/*.py.
    CO (Denver) lives on the US Redshift cluster since May 2026 and has no saved
    queries for trips / GMV / driver MAU, so those run as ad-hoc SQL (see CO_* below).
    CO Platform Revenue is the flat $1.20 per completed trip the sheet uses.

Run:
    python investor_data.py                 # previous month -> Excel -> Slack
    python investor_data.py 2026-09         # a specific month (YYYY-MM)
    python investor_data.py 2026-09 --no-slack   # keep the Excel locally, skip upload
"""

import calendar
import io
import os
import sys
import time
from datetime import datetime, timedelta

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

import pandas as pd
import requests
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from utils.helpers import Query, Redash
from utils.slack import SlackBot


# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

US_DATA_SOURCE_ID = 48          # Redash data source of the US prod cluster (NY / CO queries 7579, 7644 ...)
CO_REGION_ID = 101              # driver_locations_log.region_id for Denver
CO_PLATFORM_FEE_PER_TRIP = 1.20 # flat $ per completed trip (sheet label: "Platform Revenue (Flat $1.20)")

METRICS = [
    ("",       "Completed Trips"),
    ("",       "GMV"),
    ("",       "Platform Revenue"),
    ("Rider",  "Registered (Cumulative Sign Up)"),
    ("Rider",  "Active (MAU)"),
    ("Rider",  "Completed"),
    ("Driver", "Registered (Cumulative Sign Up)"),
    ("Driver", "Active (MAU)"),
    ("Driver", "Completed"),
]
KEYS = ["completed", "gmv", "platform_revenue",
        "rider_registered", "rider_mau", "rider_completed",
        "driver_registered", "driver_mau", "driver_completed"]


# ---------------------------------------------------------------------------
# Date helpers
# ---------------------------------------------------------------------------

def month_info(month=None):
    """Return (first, last, label) for the report month; `month` is 'YYYY-MM' or None (= previous month)."""
    if month:
        first = datetime.strptime(month + "-01", "%Y-%m-%d")
    else:
        today = datetime.today()
        first = (today.replace(day=1) - timedelta(days=1)).replace(day=1)
    last = first.replace(day=calendar.monthrange(first.year, first.month)[1])
    return first.strftime("%Y-%m-%d"), last.strftime("%Y-%m-%d"), first.strftime("%b_%Y")


def dr(start, end):
    return {"start": start, "end": end}


# ---------------------------------------------------------------------------
# Extraction helpers
# ---------------------------------------------------------------------------

def _num(v):
    try:
        if v is None or pd.isna(v):
            return None
        return float(v)
    except (TypeError, ValueError):
        return None


def first_val(df, column):
    """Row 0 of `column`, or None."""
    if df is None or not hasattr(df, "empty") or df.empty or column not in df.columns:
        return None
    return _num(df.iloc[0][column])


def month_rows(df, date_col, month):
    """Rows whose `date_col` falls in `month` (YYYY-MM). Empty frame if none."""
    if df is None or not hasattr(df, "empty") or df.empty or date_col not in df.columns:
        return pd.DataFrame()
    s = pd.to_datetime(df[date_col], errors="coerce").dt.strftime("%Y-%m")
    return df[s == month]


def month_val(df, date_col, month, column, filters=None):
    """First `column` value of the row for `month` (after optional equality filters)."""
    d = month_rows(df, date_col, month)
    for col, val in (filters or {}).items():
        if col not in d.columns:
            return None
        d = d[d[col].astype(str).str.upper() == str(val).upper()]
    if d.empty:
        return None
    return _num(d.iloc[0][column])


def month_sum(df, date_col, month, column):
    """Sum of `column` over every row for `month` (e.g. all cities / vehicle types)."""
    d = month_rows(df, date_col, month)
    if d.empty or column not in d.columns:
        return None
    return _num(pd.to_numeric(d[column], errors="coerce").sum())


def run_adhoc(sql, data_source_id):
    """Run an ad-hoc SQL string via the Redash API (used for CO, which has no saved queries)."""
    base = os.getenv("REDASH_BASE_URL").rstrip("/")
    headers = {"Authorization": f"Key {os.getenv('REDASH_API_KEY')}"}
    body = {"query": sql, "data_source_id": data_source_id, "max_age": 0}
    res = requests.post(f"{base}/api/query_results", json=body, headers=headers, timeout=300)
    res.raise_for_status()
    payload = res.json()
    if "job" in payload:
        job_id = payload["job"]["id"]
        for _ in range(400):
            job = requests.get(f"{base}/api/jobs/{job_id}", headers=headers, timeout=60).json()["job"]
            if job["status"] == 3:
                result_id = job["query_result_id"]
                break
            if job["status"] == 4:
                raise RuntimeError(f"ad-hoc query failed: {job.get('error')}")
            time.sleep(1.5)
        else:
            raise TimeoutError("ad-hoc query did not finish")
    else:
        result_id = payload["query_result"]["id"]
    csv = requests.get(f"{base}/api/query_results/{result_id}.csv", headers=headers, timeout=300)
    csv.raise_for_status()
    if not csv.text.strip():
        return pd.DataFrame()
    return pd.read_csv(io.StringIO(csv.text))


# ---------------------------------------------------------------------------
# Excel output
# ---------------------------------------------------------------------------

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
GROUP_FILL  = PatternFill("solid", fgColor="D9E1F2")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
GROUP_FONT  = Font(name="Calibri", bold=True, size=10)
BODY_FONT   = Font(name="Calibri", size=10)


def write_sheet(ws, label, values, sources):
    """values / sources: dicts keyed by KEYS. Layout mirrors the Google Sheet (group | metric | month)."""
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 52

    for c, txt in enumerate(["", "Metric", label, "Source"], 1):
        cell = ws.cell(row=1, column=c, value=txt)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center" if c == 3 else "left")

    row = 2
    for (group, name), key in zip(METRICS, KEYS):
        g = ws.cell(row=row, column=1, value=group or None)
        m = ws.cell(row=row, column=2, value=name)
        v = ws.cell(row=row, column=3, value=values.get(key))
        s = ws.cell(row=row, column=4, value=sources.get(key))
        for cell in (g, m, v, s):
            cell.font = BODY_FONT
        if group:
            g.font = GROUP_FONT
            g.fill = GROUP_FILL
        v.number_format = "#,##0.00" if key in ("gmv", "platform_revenue") else "#,##0"
        row += 1

    ws.freeze_panes = "C2"


# ---------------------------------------------------------------------------
# Shared (all markets): cumulative sign-ups
# ---------------------------------------------------------------------------

def fetch_registered(redash, first, last):
    """8095 / 7913 return one row per region for the month -> {region: value}."""
    redash.run_queries([
        Query(8095, params={"Date Range": dr(first, last)}),   # Regional - Cumulative Rider Sign Up
        Query(7913, params={"Date Range": dr(first, last)}),   # Regional - Cumulative Driver Sign Up
    ])
    month = first[:7]
    riders = month_rows(redash.get_result(8095), "month", month)
    drivers = month_rows(redash.get_result(7913), "month", month)
    r = {str(k).upper(): _num(v) for k, v in zip(riders.get("region", []), riders.get("cumulative_registered_riders", []))}
    d = {str(k).upper(): _num(v) for k, v in zip(drivers.get("region", []), drivers.get("cumulative_registered_drivers", []))}
    return r, d


REG_SRC_R = "8095 cumulative_registered_riders"
REG_SRC_D = "7913 cumulative_registered_drivers"


# ---------------------------------------------------------------------------
# Market fetchers — each returns (values, sources)
# ---------------------------------------------------------------------------

def fetch_sg(redash, first, last, reg_r, reg_d):
    month = first[:7]
    redash.run_queries([
        Query(2183, params={"date": last}),                                   # SG|RH Perf Monthly Completed and Trips
        Query(6561, params={"date_range": dr(first, last), "region": "SG"}),  # GMV Calculations
        Query(6189, params={"Date Range": dr(first, last), "region": "SG"}),  # KPI Tracking - System Fee Revenue
        Query(2187, params={"date": first}),                                  # SG|RH Perf Monthly 18-20 (rider MAU, GA4)
        Query(2194, params={"date": first}),                                  # SG|RH Perf Monthly 29-30 (rider completed)
        Query(2203, params={"date": first}),                                  # SG|RH Perf Monthly 41,43 (driver online count)
    ])
    q2183, q6561, q6189 = redash.get_result(2183), redash.get_result(6561), redash.get_result(6189)
    q2187, q2194, q2203 = redash.get_result(2187), redash.get_result(2194), redash.get_result(2203)
    values = {
        "completed":         first_val(q2183, "completed"),
        "gmv":               month_sum(q6561, "ride_month", month, "total_gmv"),
        "platform_revenue":  month_sum(q6189, "month", month, "total_system_fee"),
        "rider_registered":  reg_r.get("SG"),
        "rider_mau":         first_val(q2187, "active_users"),
        "rider_completed":   first_val(q2194, "completed_monthly"),
        "driver_registered": reg_d.get("SG"),
        "driver_mau":        first_val(q2203, "online_driver_count"),
        "driver_completed":  first_val(q2183, "completed_drivers"),
    }
    sources = {
        "completed": "2183 completed", "gmv": "6561 sum(total_gmv) all car groups",
        "platform_revenue": "6189 sum(total_system_fee)", "rider_registered": REG_SRC_R,
        "rider_mau": "2187 active_users (GA4)", "rider_completed": "2194 completed_monthly",
        "driver_registered": REG_SRC_D, "driver_mau": "2203 online_driver_count",
        "driver_completed": "2183 completed_drivers",
    }
    return values, sources


def fetch_hk(redash, first, last, reg_r, reg_d):
    month = first[:7]
    redash.run_queries([
        Query(3771, params={"date_range": dr(first, last)}),                  # HK PS | Rides
        Query(6561, params={"date_range": dr(first, last), "region": "HK"}),
        Query(6189, params={"Date Range": dr(first, last), "region": "HK"}),
        Query(3773, params={"date_range": dr(first, last)}),                  # HK PS | Active Riders (GA4)
        Query(3788, params={"date_range": dr(first, last)}),                  # HK PS | Rider Daily
        Query(3775, params={"date_range": dr(first, last)}),                  # HK PS | Active Drivers
    ])
    q3771, q6561, q6189 = redash.get_result(3771), redash.get_result(6561), redash.get_result(6189)
    q3773, q3788, q3775 = redash.get_result(3773), redash.get_result(3788), redash.get_result(3775)
    values = {
        "completed":         first_val(q3771, "completed"),
        "gmv":               month_sum(q6561, "ride_month", month, "total_gmv"),
        "platform_revenue":  month_sum(q6189, "month", month, "total_system_fee"),
        "rider_registered":  reg_r.get("HK"),
        "rider_mau":         first_val(q3773, "active_users"),
        "rider_completed":   first_val(q3788, "completed_monthly"),
        "driver_registered": reg_d.get("HK"),
        "driver_mau":        first_val(q3775, "online"),
        "driver_completed":  first_val(q3771, "completed_drivers"),
    }
    sources = {
        "completed": "3771 completed", "gmv": "6561 sum(total_gmv)",
        "platform_revenue": "6189 sum(total_system_fee)", "rider_registered": REG_SRC_R,
        "rider_mau": "3773 active_users (GA4)", "rider_completed": "3788 completed_monthly",
        "driver_registered": REG_SRC_D, "driver_mau": "3775 online",
        "driver_completed": "3771 completed_drivers",
    }
    return values, sources


def fetch_th(redash, first, last, reg_r, reg_d):
    month = first[:7]
    tz = 7
    redash.run_queries([
        Query(3106, params={"date": last}),                                   # TH PS | Rides Summary
        Query(6561, params={"date_range": dr(first, last), "region": "TH"}),
        Query(6189, params={"Date Range": dr(first, last), "region": "TH"}),
        Query(2667, params={"region": "Thailand", "timezone": tz, "date": last}),  # RH Perf Sheet - Active Riders (GA4)
        Query(3112, params={"region": "TH", "timezone": tz, "date": last}),        # TH PS | Rider Daily
        Query(2669, params={"region": 8, "timezone": tz, "date": last}),           # RH Perf Sheet - Driver Online
    ])
    q3106, q6561, q6189 = redash.get_result(3106), redash.get_result(6561), redash.get_result(6189)
    q2667, q3112, q2669 = redash.get_result(2667), redash.get_result(3112), redash.get_result(2669)
    values = {
        "completed":         first_val(q3106, "completed"),
        "gmv":               month_sum(q6561, "ride_month", month, "total_gmv"),
        "platform_revenue":  month_sum(q6189, "month", month, "total_system_fee"),
        "rider_registered":  reg_r.get("TH"),
        "rider_mau":         first_val(q2667, "active_users"),
        "rider_completed":   first_val(q3112, "completed_monthly"),
        "driver_registered": reg_d.get("TH"),
        "driver_mau":        first_val(q2669, "online_driver_count"),
        "driver_completed":  first_val(q3106, "completed_drivers"),
    }
    sources = {
        "completed": "3106 completed", "gmv": "6561 sum(total_gmv) BIKE+CAR",
        "platform_revenue": "6189 sum(total_system_fee)", "rider_registered": REG_SRC_R,
        "rider_mau": "2667 active_users (GA4)", "rider_completed": "3112 completed_monthly",
        "driver_registered": REG_SRC_D, "driver_mau": "2669 online_driver_count",
        "driver_completed": "3106 completed_drivers",
    }
    return values, sources


def fetch_kh(redash, first, last, reg_r, reg_d):
    month = first[:7]
    redash.run_queries([
        Query(1625, params={"date": first}),                                  # KH Perf sheet - All trips
        Query(6008, params={"Date Range": dr(first, last), "region": "KH"}),  # Regional - GMV Trend
        Query(6189, params={"Date Range": dr(first, last), "region": "KH"}),
        Query(5349, params={"date": first}),                                  # KH PS | MAU (GA4)
        Query(2385, params={"date": first}),                                  # KH Perf sheet - Monthly & daily Users
        Query(2373, params={"date": first}),                                  # KH Perf Sheet - Active Drivers by App Activity (GA4)
    ])
    q1625, q6008, q6189 = redash.get_result(1625), redash.get_result(6008), redash.get_result(6189)
    q5349, q2385, q2373 = redash.get_result(5349), redash.get_result(2385), redash.get_result(2373)
    values = {
        "completed":         first_val(q1625, "fin"),
        "gmv":               month_val(q6008, "trip_month", month, "gmv"),
        "platform_revenue":  month_sum(q6189, "month", month, "total_system_fee"),
        "rider_registered":  reg_r.get("KH"),
        "rider_mau":         first_val(q5349, "active_users"),
        "rider_completed":   first_val(q2385, "finished_rider_count"),
        "driver_registered": reg_d.get("KH"),
        "driver_mau":        first_val(q2373, "mau_driver"),
        "driver_completed":  first_val(q1625, "completed_drivers"),
    }
    sources = {
        "completed": "1625 fin", "gmv": "6008 gmv",
        "platform_revenue": "6189 sum(total_system_fee) PNH + KH-OTHERS", "rider_registered": REG_SRC_R,
        "rider_mau": "5349 active_users (GA4)", "rider_completed": "2385 finished_rider_count",
        "driver_registered": REG_SRC_D, "driver_mau": "2373 mau_driver (GA4)",
        "driver_completed": "1625 completed_drivers",
    }
    return values, sources


def fetch_vn(redash, first, last, reg_r, reg_d):
    month = first[:7]
    redash.run_queries([
        Query(4562, params={"date_range": dr(first, last), "city": "ALL"}),  # VN PS | Rides
        Query(6008, params={"Date Range": dr(first, last), "region": "VN"}),
        Query(6189, params={"Date Range": dr(first, last), "region": "VN"}),
        Query(4565, params={"date_range": dr(first, last), "city": "ALL"}),  # VN PS | Active Riders (GA4)
        Query(4581, params={"date_range": dr(first, last), "city": "ALL"}),  # VN PS | Rider Daily
        Query(4568, params={"date_range": dr(first, last), "city": "ALL"}),  # VN PS | Active Drivers
    ])
    q4562, q6008, q6189 = redash.get_result(4562), redash.get_result(6008), redash.get_result(6189)
    q4565, q4581, q4568 = redash.get_result(4565), redash.get_result(4581), redash.get_result(4568)
    values = {
        "completed":         first_val(q4562, "completed"),
        "gmv":               month_val(q6008, "trip_month", month, "gmv"),
        "platform_revenue":  month_sum(q6189, "month", month, "total_system_fee"),
        "rider_registered":  reg_r.get("VN"),
        "rider_mau":         first_val(q4565, "active_users"),
        "rider_completed":   first_val(q4581, "completed_monthly_all"),
        "driver_registered": reg_d.get("VN"),
        "driver_mau":        first_val(q4568, "online"),
        "driver_completed":  first_val(q4562, "completed_drivers"),
    }
    sources = {
        "completed": "4562 completed (city=ALL)", "gmv": "6008 gmv",
        "platform_revenue": "6189 sum(total_system_fee) HCM + HAN, 2W + 4W", "rider_registered": REG_SRC_R,
        "rider_mau": "4565 active_users (GA4, city=ALL)", "rider_completed": "4581 completed_monthly_all",
        "driver_registered": REG_SRC_D, "driver_mau": "4568 online (city=ALL)",
        "driver_completed": "4562 completed_drivers",
    }
    return values, sources


def fetch_ny(redash, first, last, reg_r, reg_d):
    # NY data lags ~1 day; run on/after the 2nd of the month for a full month.
    month = first[:7]
    redash.run_queries([
        Query(7579, params={"date": first}),                     # NY PS - Perf Monthly Unique
        Query(7644, params={"date": first}),                     # Monthly Trips and GMV - NY (month-START date)
        Query(7655, params={"Date Range": dr(first, last)}),     # KPI Tracking - NY System Fee Revenue
        Query(7621, params={"date": first}),                     # NY PS | Active Users (GA4)
        Query(7586, params={"date": first}),                     # NY PS | Drivers Online and Utilisation Hours
    ])
    q7579, q7644, q7655 = redash.get_result(7579), redash.get_result(7644), redash.get_result(7655)
    q7621, q7586 = redash.get_result(7621), redash.get_result(7586)
    values = {
        "completed":         first_val(q7579, "nyc_completed"),
        "gmv":               first_val(q7644, "gmv"),
        "platform_revenue":  month_sum(q7655, "month", month, "total_system_fee"),
        "rider_registered":  reg_r.get("NY"),
        "rider_mau":         first_val(q7621, "active_users"),
        "rider_completed":   first_val(q7579, "nyc_rider_completed"),
        "driver_registered": reg_d.get("NY"),
        "driver_mau":        first_val(q7586, "online_driver_count"),
        "driver_completed":  first_val(q7579, "nyc_driver_completed"),
    }
    sources = {
        "completed": "7579 nyc_completed", "gmv": "7644 gmv",
        "platform_revenue": "7655 sum(total_system_fee)", "rider_registered": REG_SRC_R,
        "rider_mau": "7621 active_users (GA4)", "rider_completed": "7579 nyc_rider_completed",
        "driver_registered": REG_SRC_D, "driver_mau": "7586 online_driver_count",
        "driver_completed": "7579 nyc_driver_completed",
    }
    return values, sources


# --- CO (Denver): ad-hoc SQL on the US cluster -------------------------------
# Completed trips are counted on the reservation-aware Denver month (like 7644);
# GMV, unique riders and unique drivers on the create-time Denver month (like 6561).
# Both bases reproduce the sheet exactly for Jun-Sep 2026.

CO_TRIPS_SQL = """
WITH base AS (
    SELECT
        date_trunc('month', convert_timezone('UTC', 'America/Denver', e.create_time::timestamp))::date AS m_create,
        date_trunc('month', convert_timezone('UTC', 'America/Denver',
            COALESCE(e.reservation_ride_start_time::timestamp, e.create_time::timestamp)))::date AS m_resv,
        e.id, e.rider_uuid, e.driver_uuid,
        COALESCE(e.ride_price::float, 0) + COALESCE(e.toll_fee::float, 0) + COALESCE(e.etc_fee::float, 0)
          + COALESCE(e.rider_system_fee::float, 0) + COALESCE(e.rider_application_fee::float, 0)
          + COALESCE(e.ride_option_fee::float, 0) + COALESCE(e.reservation_fee::float, 0) AS gmv_amount
    FROM tada_ride_service_us.ride_entity e
    WHERE e.region = 'CO'
      AND e.rider_uuid IS NOT NULL
      AND e.ride_status = 70
      AND e.create_time::timestamp >= DATEADD(day, -3, DATE '{first}')
      AND e.create_time::timestamp <  DATEADD(day,  3, DATE '{last}')
)
SELECT
    (SELECT COUNT(*) FROM base WHERE m_resv = DATE '{first}')                        AS completed,
    (SELECT SUM(gmv_amount) FROM base WHERE m_create = DATE '{first}')               AS gmv,
    (SELECT COUNT(DISTINCT rider_uuid) FROM base WHERE m_create = DATE '{first}')    AS completed_riders,
    (SELECT COUNT(DISTINCT driver_uuid) FROM base WHERE m_create = DATE '{first}')   AS completed_drivers
"""

CO_DRIVER_MAU_SQL = """
SELECT COUNT(DISTINCT l.driver_uuid) AS online_driver_count
FROM tada_bq.driver_locations_log l
WHERE l.region_id = {region_id}
  AND l.car_type IN (0,1,4,11,12,13,21,22,23,31,32,10000,10001,10002,10003)
  AND l.recorded_at::timestamp >= DATEADD(day, -1, DATE '{first}')
  AND l.recorded_at::timestamp <  DATEADD(day,  2, DATE '{last}')
  AND date_trunc('month', convert_timezone('UTC', 'America/Denver', l.recorded_at::timestamp))::date = DATE '{first}'
"""


def fetch_co(redash, first, last, reg_r, reg_d):
    month = first[:7]
    redash.run_queries([
        Query(8096, params={"Date Range": dr(first, last)}),     # CO - Active Rider/User (GA4, Denver)
    ])
    q8096 = redash.get_result(8096)
    trips = run_adhoc(CO_TRIPS_SQL.format(first=first, last=last), US_DATA_SOURCE_ID)
    drv = run_adhoc(CO_DRIVER_MAU_SQL.format(first=first, last=last, region_id=CO_REGION_ID), US_DATA_SOURCE_ID)
    completed = first_val(trips, "completed")
    values = {
        "completed":         completed,
        "gmv":               first_val(trips, "gmv"),
        "platform_revenue":  (completed * CO_PLATFORM_FEE_PER_TRIP) if completed is not None else None,
        "rider_registered":  reg_r.get("CO"),
        "rider_mau":         month_val(q8096, "month", month, "active_riders"),
        "rider_completed":   first_val(trips, "completed_riders"),
        "driver_registered": reg_d.get("CO"),
        "driver_mau":        first_val(drv, "online_driver_count"),
        "driver_completed":  first_val(trips, "completed_drivers"),
    }
    sources = {
        "completed": "ad-hoc US cluster: completed rides, reservation-aware Denver month",
        "gmv": "ad-hoc US cluster: GMV (6561 formula), create-time Denver month",
        "platform_revenue": f"completed x {CO_PLATFORM_FEE_PER_TRIP:.2f} (flat fee)",
        "rider_registered": REG_SRC_R,
        "rider_mau": "8096 active_riders (GA4, Denver)",
        "rider_completed": "ad-hoc US cluster: distinct riders with a completed ride",
        "driver_registered": REG_SRC_D,
        "driver_mau": f"ad-hoc US cluster: driver_locations_log region_id {CO_REGION_ID}, car-type filtered",
        "driver_completed": "ad-hoc US cluster: distinct drivers with a completed ride",
    }
    return values, sources


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    load_dotenv()

    redash = Redash(
        key=os.getenv("REDASH_API_KEY"),
        base_url=os.getenv("REDASH_BASE_URL"),
    )

    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    flags = {a for a in sys.argv[1:] if a.startswith("--")}
    month_arg = args[0] if args else None                        # optional YYYY-MM
    upload = "--no-slack" not in flags

    first, last, label = month_info(month_arg)

    print(f"Investor Data - {label}")
    print(f"  Period : {first} -> {last}\n")

    print("-> cumulative sign-ups (8095 / 7913)")
    reg_r, reg_d = fetch_registered(redash, first, last)

    tasks = [
        ("SG",  fetch_sg),
        ("VN",  fetch_vn),
        ("KH",  fetch_kh),
        ("TH",  fetch_th),
        ("HK",  fetch_hk),
        ("NYC", fetch_ny),
        ("CO",  fetch_co),
    ]

    wb = Workbook()
    wb.remove(wb.active)

    for name, fn in tasks:
        print(f"-> {name}")
        ws = wb.create_sheet(name)
        try:
            values, sources = fn(redash, first, last, reg_r, reg_d)
            write_sheet(ws, label, values, sources)
            missing = [k for k in KEYS if values.get(k) is None]
            print("   done" + (f"  (no value: {', '.join(missing)})" if missing else ""))
        except Exception as exc:
            import traceback
            print(f"   FAILED: {exc}")
            traceback.print_exc()
            ws["A1"] = f"FAILED: {exc}"

    output_file = f"Investor_Data_{label}.xlsx"
    wb.save(output_file)
    print(f"\nSaved: {output_file}")

    if not upload:
        print("Slack upload skipped (--no-slack)")
        return

    slack = SlackBot()
    slack.uploadFile(
        output_file,
        os.getenv("SLACK_CHANNEL"),
        f"Investor Data for {label.replace('_', ' ')}",
    )
    print("Uploaded to Slack")

    try:
        os.remove(output_file)
    except OSError:
        pass


if __name__ == "__main__":
    main()
